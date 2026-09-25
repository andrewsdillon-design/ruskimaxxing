"""Build the standalone Excel workbook.

Every weight is an Excel formula tied to the lifts entered on the "Start Here"
sheet, so the file works on its own in Excel, LibreOffice or Google Sheets.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from ruskimaxxing.program import MAIN_LIFTS, build_program

TITLE = "RuskiMaxxing - 12-Week Beginner Strength & Mass"
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")
BAND_FILL = PatternFill("solid", fgColor="EEF3F8")
PHASE_FILLS = {
    "Accumulation": "DCEFDC", "Transmutation": "FCE9D2", "Realization": "F8D7D7",
    "Deload": "E6E6E6", "Test": "DDE3F7",
}
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

INSTRUCTIONS = [
    "How to use this sheet",
    "1. Pick lb or kg and your plate rounding (5 lb / 2.5 kg is typical).",
    "2. For each lift, enter a recent set: the weight and how many clean reps you did (1-10 reps).",
    "   If you know your true 1-rep max, enter it with reps = 1.",
    "   New lifter? Work up to a set of 5 you could just barely do, and enter that.",
    "3. Open the Program sheet. Every weight fills in automatically.",
    "4. Train 3 days a week with a rest day between (e.g. Mon / Wed / Fri).",
    "5. Warm up before every main lift: empty bar x 10, then 3-4 lighter sets working up.",
    "6. Log each session in the Done / Actual reps / Notes columns.",
    "7. Week 12: enter your AMRAP reps - the sheet estimates your new 1RMs.",
    "   Put those numbers here and run the 12 weeks again.",
    "",
    "Eat enough to grow: a small calorie surplus and roughly 0.7-1 g protein per lb bodyweight (1.6-2.2 g/kg).",
    "Sleep 7-9 hours. Progress comes from showing up consistently.",
    "",
    "Free for everyone - MIT license. Not medical advice; check with a doctor before starting a new program.",
]


def _style_header(cells):
    for c in cells:
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _start_sheet(ws, inputs: dict | None):
    ws.title = "Start Here"
    ws["A1"] = TITLE
    ws["A1"].font = Font(bold=True, size=16)
    ws["A3"], ws["B3"] = "Units", (inputs or {}).get("units", "lb")
    ws["A4"], ws["B4"] = "Round weights to nearest", (inputs or {}).get("increment", 5)
    for ref in ("B3", "B4"):
        ws[ref].fill, ws[ref].border = INPUT_FILL, BOX
    units = DataValidation(type="list", formula1='"lb,kg"', allow_blank=False)
    ws.add_data_validation(units)
    units.add("B3")

    ws.append([])
    ws.append(["Lift", "Weight lifted", "Reps done", "Estimated 1RM"])
    _style_header(ws[6])
    lift_refs = {}
    maxes = (inputs or {}).get("lifts", {})
    for i, lift in enumerate(MAIN_LIFTS):
        r = 7 + i
        weight, reps = maxes.get(lift, (None, None))
        ws.cell(r, 1, lift).border = BOX
        for col, val in ((2, weight), (3, reps)):
            c = ws.cell(r, col, val)
            c.fill, c.border = INPUT_FILL, BOX
        c = ws.cell(r, 4, f'=IF(OR(B{r}="",C{r}=""),"",IF(C{r}=1,B{r},ROUND(B{r}*(1+C{r}/30),1)))')
        c.border, c.font = BOX, Font(bold=True)
        lift_refs[lift] = f"'Start Here'!$D${r}"

    for i, line in enumerate(INSTRUCTIONS):
        c = ws.cell(13 + i, 1, line)
        if i == 0:
            c.font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 26
    for col in "BCD":
        ws.column_dimensions[col].width = 16
    return lift_refs


def _program_sheet(ws, lift_refs: dict):
    ws.title = "Program"
    headers = ["Week", "Phase", "Day", "Exercise", "Sets", "Reps", "% 1RM", "Weight",
               "Done", "Actual reps", "Est. new 1RM", "Notes / how it felt", "Guidance"]
    ws.append(headers)
    _style_header(ws[1])
    ws.freeze_panes = "E2"

    done = DataValidation(type="list", formula1='"✓,✗"', allow_blank=True)
    ws.add_data_validation(done)

    row = 2
    for session in build_program():
        phase_fill = PatternFill("solid", fgColor=PHASE_FILLS[session.phase])
        for p in session.exercises:
            ref = lift_refs.get(p.exercise)
            weight = (f'=IF({ref}="","",MROUND({ref}*G{row}/100,\'Start Here\'!$B$4))'
                      if p.is_main and ref else None)
            new_max = (f'=IF(OR(J{row}="",H{row}=""),"",ROUND(H{row}*(1+J{row}/30),0))'
                       if p.reps == "AMRAP" else None)
            values = [session.week, session.phase, session.day, p.exercise,
                      p.sets or None, p.reps, p.percent, weight, None, None, new_max, None, p.note]
            for col, val in enumerate(values, start=1):
                c = ws.cell(row, col, val)
                c.border = BOX
                c.alignment = Alignment(vertical="center", wrap_text=col == 13)
            ws.cell(row, 2).fill = phase_fill
            ws.cell(row, 8).font = Font(bold=True)
            for col in (9, 10, 12):
                ws.cell(row, col).fill = INPUT_FILL
            if p.is_main:
                ws.cell(row, 4).font = Font(bold=True)
            done.add(f"I{row}")
            row += 1

    last = row - 1
    # Shade alternate weeks so each week reads as a group
    ws.conditional_formatting.add(
        f"A2:H{last}", FormulaRule(formula=["MOD($A2,2)=0"], fill=BAND_FILL))
    widths = [7, 14, 16, 26, 6, 8, 8, 10, 7, 11, 12, 28, 48]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.auto_filter.ref = f"A1:M{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def build_workbook(path: str | Path, inputs: dict | None = None) -> Path:
    """Write the workbook to `path`.

    `inputs` optionally pre-fills the Start Here sheet:
    {"units": "lb", "increment": 5, "lifts": {"Squat": (weight, reps), ...}}
    """
    wb = Workbook()
    lift_refs = _start_sheet(wb.active, inputs)
    _program_sheet(wb.create_sheet(), lift_refs)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
