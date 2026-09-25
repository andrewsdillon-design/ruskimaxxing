import csv
import shutil
import subprocess

import pytest
from openpyxl import load_workbook

from ruskimaxxing.excel import build_workbook
from ruskimaxxing.program import build_program

INPUTS = {"units": "lb", "increment": 5,
          "lifts": {"Squat": (250, 5), "Bench Press": (185, 1), "Deadlift": (300, 3), "Overhead Press": (95, 8)}}


def test_workbook_structure(tmp_path):
    path = build_workbook(tmp_path / "p.xlsx", INPUTS)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Start Here", "Program"]
    ws = wb["Program"]
    rows = ws.max_row - 1
    assert rows == sum(len(s.exercises) for s in build_program())
    assert ws["H2"].value.startswith("=IF(") and "MROUND" in ws["H2"].value
    assert wb["Start Here"]["B7"].value == 250


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_formulas_calculate_in_libreoffice(tmp_path):
    build_workbook(tmp_path / "p.xlsx", INPUTS)
    subprocess.run(["soffice", "--headless", "--convert-to", "csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,false,false,false,-1",
                    "--outdir", str(tmp_path), str(tmp_path / "p.xlsx")],
                   check=True, capture_output=True, timeout=120)
    with open(tmp_path / "p-Program.csv", newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    first = rows[0]
    assert first["Exercise"] == "Squat" and first["% 1RM"] == "70"
    # 250 x 5 -> Epley 291.7 -> 70% = 204.2 -> nearest 5 = 205
    assert first["Weight"] == "205"
    bench = next(r for r in rows if r["Exercise"] == "Bench Press")
    assert bench["Weight"] == "130"  # 185 * 0.70 = 129.5 -> 130
